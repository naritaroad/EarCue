#!/usr/bin/env python3
"""Build the fictional "Green Streets" demo: a bilingual transcript (.docx), a slide deck (.pptx)
and the expected EarCue cue file for presenter P2.

Usage: python tools/demo/build_demo.py   (needs python-docx and python-pptx)
Output: examples/green-streets/
"""
import json
import pathlib
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt as PPt

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import content as c  # noqa: E402

OUT = pathlib.Path(__file__).resolve().parents[2] / "examples" / "green-streets"

AMBER = (0xE4, 0x8A, 0x12)
INK = (0x1B, 0x1C, 0x20)
DIM = (0x5D, 0x62, 0x70)
PAPER = (0xF6, 0xF3, 0xEE)
CARD = (0xFF, 0xFF, 0xFF)


def minutes_en(m):
    whole, half = int(m), m - int(m) >= 0.5
    return f"{whole} min" if not half else f"{whole}.5 min"


def minutes_zh(m):
    whole, half = int(m), m - int(m) >= 0.5
    return f"{whole} 分" + (" 30 秒" if half else "")


def speakers(slide):
    return " × ".join(c.who(s) for s in slide["speakers"])


# ---------------------------------------------------------------- transcript

def build_docx(path):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def para(text, bold=False, size=None, color=None, align=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*color)
        if align:
            p.alignment = align
        return p

    para(f"{c.TITLE_ZH} 共同報告逐字稿 V1", bold=True, size=18)
    para(f"{c.TITLE} · Co-presenting Transcript V1 (English / Chinese)", bold=True, size=14)
    para(c.CITATION, color=DIM)
    para("Roles: " + c.ROLES)
    para(c.ROLES_ZH)
    para("Every line is in English first, then Chinese. Slide numbers and English titles come from the demo deck.")
    para("每一句都是英文在前、中文在後；頁碼與英文標題對應示範投影片。")
    para("Fictional demo: the paper, the data and the presenters are made up for testing EarCue.", color=AMBER)

    para("Overview · 總覽", bold=True, size=14)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for cell, h in zip(table.rows[0].cells, ["Slide 頁", "Part 部分", "Title 標題", "Speaker 負責", "Time 時間"]):
        cell.text = h
    for s in c.SLIDES:
        row = table.add_row().cells
        row[0].text = s["id"]
        row[1].text = f"{s['part']}\n{s['part_zh']}"
        row[2].text = f"{s['title']}\n{s['title_zh']}"
        row[3].text = speakers(s)
        row[4].text = f"{minutes_en(s['minutes'])}\n{minutes_zh(s['minutes'])}"
    total = sum(s["minutes"] for s in c.SLIDES)
    para(f"Total (Q&A not included) · 合計（不含 Q&A）: {total:g} min")

    for s in c.SLIDES:
        doc.add_paragraph()
        para(f"{s['id']} | {s['title']}", bold=True, size=14)
        para(f"{s['id']}｜{s['title_zh']}", bold=True, size=12)
        para(f"Speaker · 負責: {speakers(s)}　　Time: about {minutes_en(s['minutes'])}　　時間：約 {minutes_zh(s['minutes'])}", color=DIM)
        para(f"Core question: {s['core'][0]}")
        para(f"核心問題：{s['core'][1]}")
        para("Slide key points · 投影片重點:", bold=True)
        for i, (en, zh) in enumerate(s["points"], 1):
            para(f"{i}. {en}")
            para(zh)
        if s["handover"]:
            code, en, zh = s["handover"]
            para(f"Hand-over question ({c.who(code)}): {en}", color=DIM)
            para(f"轉場提問（{c.who(code)}）：{zh}", color=DIM)
        if s["mode"] == "two":
            para("Script (two-person highlight) · 逐字稿（雙人亮點：一問一答）", bold=True)
        else:
            para("Script (one person) · 逐字稿（一個人講完）", bold=True)
        for code, en, zh in s["lines"]:
            para(f"{c.who(code)}: {en}")
            para(f"{c.who(code)}：{zh}")
    doc.save(path)


# ---------------------------------------------------------------- slides

def build_pptx(path):
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]
    W, H = prs.slide_width, prs.slide_height

    def box(slide, x, y, w, h, fill, radius=False):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = PRGB(*fill)
        shape.line.fill.background()
        if radius:
            shape.adjustments[0] = 0.12
        return shape

    def text(slide, x, y, w, h, value, size, color, bold=False):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        run = tf.paragraphs[0].add_run()
        run.text = value
        run.font.size = PPt(size)
        run.font.bold = bold
        run.font.color.rgb = PRGB(*color)
        return tb

    margin = Emu(700000)
    for s in c.SLIDES:
        slide = prs.slides.add_slide(blank)
        box(slide, 0, 0, W, H, PAPER)
        box(slide, 0, 0, Emu(120000), H, AMBER)
        text(slide, margin, Emu(520000), W - 2 * margin, Emu(400000), s["kicker"], 14, AMBER, bold=True)
        if s["id"] == "P1":
            text(slide, margin, Emu(1500000), W - 2 * margin, Emu(1200000), c.TITLE, 54, INK, bold=True)
            text(slide, margin, Emu(2800000), W - 2 * margin, Emu(500000), c.SUBTITLE, 16, DIM, bold=True)
            text(slide, margin, Emu(3400000), W - 2 * margin, Emu(600000), c.TAGLINE, 26, INK)
            text(slide, margin, Emu(5600000), W - 2 * margin, Emu(500000), "  ·  ".join(c.who(k) for k in c.PEOPLE), 13, DIM)
        else:
            text(slide, margin, Emu(950000), W - 2 * margin, Emu(1100000), s["title"], 38, INK, bold=True)
            text(slide, margin, Emu(2050000), W - 2 * margin, Emu(450000), speakers(s), 16, DIM, bold=True)
            n = len(s["points"])
            gap = Emu(250000)
            cw = int((W - 2 * margin - gap * (n - 1)) / n)
            for i, (en, _) in enumerate(s["points"]):
                x = margin + i * (cw + gap)
                box(slide, x, Emu(2900000), cw, Emu(2300000), CARD, radius=True)
                text(slide, x + Emu(250000), Emu(3100000), cw - Emu(500000), Emu(500000), f"0{i + 1}", 20, AMBER, bold=True)
                text(slide, x + Emu(250000), Emu(3650000), cw - Emu(500000), Emu(1400000), en, 20, INK)
        text(slide, margin, H - Emu(560000), W - 2 * margin, Emu(400000), f"{s['id']}  ·  Fictional demo for EarCue", 11, DIM)
        notes = slide.notes_slide.notes_text_frame
        notes.text = "\n".join(f"{c.who(code)}: {en}" for code, en, _ in s["lines"])
    prs.save(path)


# ---------------------------------------------------------------- expected cue file

def build_cues(path, role="P2"):
    sections = []
    for s in c.SLIDES:
        if role not in s["speakers"]:
            continue
        cues = []
        if s["handover"] and s["mode"] == "one" and s["handover"][0] != role:
            code, en, zh = s["handover"]
            cues.append({"partner": c.who(code), "text": en, "note": "Hand-over question"})
        for code, en, zh in s["lines"]:
            if code == role:
                cues.append({"text": en, "note": zh})
            else:
                cues.append({"partner": c.who(code), "text": en})
        sections.append({"title": s["title"], "slide": s["id"], "minutes": s["minutes"], "cues": cues})
    doc = {
        "format": "earcue/1",
        "title": f"{c.TITLE} · {role} Background",
        "speaker": c.who(role),
        "lang": "en-US",
        "tags": ["demo", "paper-review", "urban-heat", "co-presenting"],
        "source": "Green Streets Transcript (Bilingual).docx",
        "sections": sections,
    }
    path.write_text(json.dumps(doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx(OUT / "Green Streets Transcript (Bilingual).docx")
    build_pptx(OUT / "Green Streets Co-presenting Deck.pptx")
    build_cues(OUT / "green-streets-p2.earcue.json")
    print("wrote", *sorted(p.name for p in OUT.iterdir()), sep="\n  ")


if __name__ == "__main__":
    main()
